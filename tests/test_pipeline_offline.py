"""Corrida completa sin internet sobre una copia del proyecto (usa las respuestas guardadas en fixtures)."""
import io
import shutil
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl import Workbook

from monitor.pipeline import run

ROOT = Path(__file__).resolve().parents[1]


def _dane_emc_file() -> bytes:
    """Anexo tipo DANE con dos filas candidatas: la nominal (no calza) y la real (calza con el histórico)."""
    hist = openpyxl.load_workbook(ROOT / "excel" / "Monitor_Retail_LATAM_original.xlsx", data_only=True)["2_Ventas_Retail"]
    real = {}
    for r in range(24, 85):
        d, v = hist.cell(r, 2).value, hist.cell(r, 5).value
        if d is not None and isinstance(v, (int, float)):
            real[(d.year, d.month)] = v * 100
    months = sorted(real) + [(2026, 3), (2026, 4)]
    real[(2026, 3)], real[(2026, 4)] = 9.5, 8.7
    wb = Workbook()
    ws = wb.active
    ws.title = "1.1"
    ws.append(["Encuesta Mensual de Comercio - variación anual (%)"])
    ws.append(["Concepto"] + [f"{['Ene','Feb','Mar','Abr','May','Jun','Jul','Ago','Sep','Oct','Nov','Dic'][m-1]}-{str(y)[2:]}" for y, m in months])
    ws.append(["Total comercio minorista - ventas nominales"] + [real[k] + 9 for k in months])
    ws.append(["Total comercio minorista - ventas reales"] + [real[k] for k in months])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_corrida_offline(tmp_path):
    for d in ("config", "src"):
        shutil.copytree(ROOT / d, tmp_path / d)
    (tmp_path / "excel").mkdir()
    shutil.copy(ROOT / "excel" / "Monitor_Retail_LATAM_original.xlsx", tmp_path / "excel" / "Monitor_Retail_LATAM_original.xlsx")
    latest = tmp_path / "data" / "raw" / "latest"
    shutil.copytree(ROOT / "tests" / "fixtures" / "latest", latest)
    (latest / "dane_emc_pagina.html").write_text(
        '<a href="/files/operaciones/EMC/anex-EMC-ComercioMinorista-abr2026.xlsx">Anexos</a>', encoding="utf-8")
    (latest / "dane_emc_archivo.xlsx").write_bytes(_dane_emc_file())
    rc = run(tmp_path, offline=True, solo=["2_Ventas_Retail", "5_Inflacion"], sin_noticias=True)
    assert rc == 0
    wb = openpyxl.load_workbook(tmp_path / "excel" / "Monitor_Retail_LATAM.xlsx", data_only=True)
    ws = wb["2_Ventas_Retail"]
    vals = {ws.cell(r, 2).value.strftime("%Y-%m"): ws.cell(r, 5).value for r in range(24, 95)
            if hasattr(ws.cell(r, 2).value, "strftime")}
    assert abs(vals["2026-04"] - 0.087) < 1e-9          # tomó la fila REAL, no la nominal
    log = wb["Log_Actualizacion"]
    textos = " ".join(str(c.value) for row in log.iter_rows() for c in row if c.value)
    assert "ventas reales" in textos
    # inflación: BIS como respaldo cuando la fuente nacional no está disponible
    wi = wb["5_Inflacion"]
    fechas = {wi.cell(r, 2).value.strftime("%Y-%m"): r for r in range(27, 100) if hasattr(wi.cell(r, 2).value, "strftime")}
    assert abs(wi.cell(fechas["2026-08"], 4).value - 0.041321553) < 1e-9
    assert (tmp_path / "output" / "respaldos").exists()
